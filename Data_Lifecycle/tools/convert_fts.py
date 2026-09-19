#!/usr/bin/env python3
"""Convert a DMBOK State Transition metadata workbook into the canonical FTS
metamodel JSON that the skill populates and the application reads.
This is the single source of truth format: states (nested to any depth),
transitions with guards, the four-predicate contracts, KA interactions,
governance chain and relationships."""
import openpyxl, json, sys, datetime, zoneinfo

SRC = sys.argv[1] if len(sys.argv) > 1 else "DMBOK_Plan_Design_State_Transition_Metadata_v2.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "plan_design.fts.json"

wb = openpyxl.load_workbook(SRC, data_only=True)

def rows(sheet):
    """Yield dict rows keyed by header for a sheet, skipping blank rows."""
    if sheet not in wb.sheetnames:
        return
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    # find header row = first row with >1 non-empty cell
    header = None
    for r in it:
        nonempty = [c for c in r if c not in (None, "")]
        if len(nonempty) > 1:
            header = [(str(c).strip() if c is not None else "") for c in r]
            break
    if not header:
        return
    for r in it:
        if all(c in (None, "") for c in r):
            continue
        d = {}
        for h, v in zip(header, r):
            if h:
                d[h] = v.strip() if isinstance(v, str) else v
        # skip rows that are entirely empty under headers
        if any(v not in (None, "") for v in d.values()):
            yield d

def g(d, *keys, default=None):
    for k in keys:
        if k in d and d[k] not in (None, ""):
            return d[k]
    return default

def split_ids(v):
    if not v:
        return []
    return [x.strip() for x in str(v).replace(",", ";").split(";") if x.strip()]

model = {}

model["meta"] = {
    "modelId": "GDS-PLAN-DESIGN",
    "name": "Plan and Design Data State Transition",
    "level": "worked-example",
    "knowledgeArea": None,
    "version": "2.0",
    "status": "Proposed / illustrative",
    "buildStamp": datetime.datetime.now(zoneinfo.ZoneInfo("Africa/Johannesburg")).strftime("%d %b %Y, %H:%M SAST"),
    "source": SRC,
    "note": "Converted from the v2 metadata workbook. Retained-decision rules are normative for this workbook only, not DMBOK v3 Editorial Board approvals."
}

# --- Rules (the 24 Important Retained Decisions) ---
model["rules"] = [{
    "id": g(d, "Rule ID"), "shortName": g(d, "Short Name"),
    "statement": g(d, "Rule Statement"), "rationale": g(d, "Rationale"),
    "appliesTo": g(d, "Applies To"), "test": g(d, "Validation Test"),
    "compliant": g(d, "Compliant Example"), "nonCompliant": g(d, "Non-Compliant Example"),
    "status": g(d, "Status"), "authority": g(d, "Authority / Scope"), "notes": g(d, "Notes"),
} for d in rows("Important Retained Decisions") if g(d, "Rule ID")]

# --- Global States ---
model["globalStates"] = [{
    "id": g(d, "State ID"), "name": g(d, "Name"), "kind": "global",
    "definition": g(d, "Definition / Purpose"), "stateType": g(d, "State Type"),
    "initial": str(g(d, "Initial?", default="")).lower().startswith("y"),
    "terminal": str(g(d, "Terminal?", default="")).lower().startswith("y"),
    "status": g(d, "Status"), "semanticClass": g(d, "Semantic Classification"),
    "nameQA": {"status": g(d, "Name QA Status"), "rationale": g(d, "Name QA Rationale")},
    "retainedRuleIds": split_ids(g(d, "Retained Rule IDs")),
} for d in rows("Global States") if g(d, "State ID")]

# --- Sub-States (parent may be a global state or another sub-state = nesting) ---
model["subStates"] = [{
    "id": g(d, "Sub-State ID"), "name": g(d, "Name"), "kind": "sub",
    "parent": g(d, "Parent State ID"), "definition": g(d, "Definition / Purpose"),
    "sequence": g(d, "Sequence Hint"), "requirement": g(d, "Requirement"),
    "readiness": str(g(d, "Readiness State?", default="")).lower().startswith("y"),
    "status": g(d, "Status"), "semanticClass": g(d, "Semantic Classification"),
    "nameQA": {"status": g(d, "Name QA Status"), "rationale": g(d, "Name QA Rationale")},
    "retainedRuleIds": split_ids(g(d, "Retained Rule IDs")),
    "notes": g(d, "Notes"),
} for d in rows("Sub-States") if g(d, "Sub-State ID")]

# --- Activities ---
model["activities"] = [{
    "id": g(d, "Activity ID"), "name": g(d, "Activity Name"),
    "permittedIn": g(d, "Permitted In Element ID"), "activityType": g(d, "Activity Type"),
    "nounVerb": g(d, "Noun / Verb"), "permissibility": g(d, "Permissibility"), "notes": g(d, "Notes"),
} for d in rows("Activities") if g(d, "Activity ID")]

# --- Transitions ---
model["transitions"] = [{
    "id": g(d, "Transition ID"), "name": g(d, "Name"), "level": g(d, "Level"),
    "source": g(d, "Source Element ID"), "target": g(d, "Target Element ID"),
    "event": g(d, "Trigger Event ID"), "transitionType": g(d, "Transition Type"),
    "optionality": g(d, "Optionality"), "guardSummary": g(d, "Guard / Rationale Summary"),
    "status": g(d, "Status"),
} for d in rows("Transitions") if g(d, "Transition ID")]

# --- Events ---
model["events"] = [{
    "id": g(d, "Event ID"), "name": g(d, "Name"), "eventType": g(d, "Event Type"),
    "meaning": g(d, "Definition / Trigger Meaning"),
} for d in rows("Events") if g(d, "Event ID")]

# --- Entry Conditions ---
model["entryConditions"] = [{
    "id": g(d, "Condition ID"), "appliesTo": g(d, "Applies To Element ID"),
    "predicate": g(d, "Predicate / Condition"), "requirement": g(d, "Requirement"),
    "conditionType": g(d, "Condition Type"),
} for d in rows("Entry Conditions") if g(d, "Condition ID")]

# --- Exit Conditions ---
model["exitConditions"] = [{
    "id": g(d, "Condition ID"), "appliesTo": g(d, "Applies To Element ID"),
    "predicate": g(d, "Predicate / Condition"), "requirement": g(d, "Requirement"),
} for d in rows("Exit Conditions") if g(d, "Condition ID")]

# --- Invariants ---
model["invariants"] = [{
    "id": g(d, "Invariant ID"), "appliesTo": g(d, "Applies To Element ID"),
    "predicate": g(d, "Invariant Predicate"), "severity": g(d, "Severity"),
} for d in rows("Invariants") if g(d, "Invariant ID")]

# --- Transition Guards ---
model["guards"] = [{
    "id": g(d, "Guard ID"), "name": g(d, "Name"), "transition": g(d, "Transition ID"),
    "predicate": g(d, "Predicate / Guard"), "scope": g(d, "Contributing Scope"),
    "requirement": g(d, "Requirement"),
} for d in rows("Transition Guards") if g(d, "Guard ID")]

# --- KA Interactions ---
model["kaInteractions"] = [{
    "id": g(d, "Interaction ID"), "subState": g(d, "Global Sub-State ID"),
    "knowledgeArea": g(d, "Knowledge Area"), "interactionType": g(d, "Interaction Type"),
    "applicability": g(d, "Applicability"), "description": g(d, "Interaction Description"),
} for d in rows("KA Interactions") if g(d, "Interaction ID")]

# --- KA Matrix ---
ka_matrix = []
for d in rows("KA Matrix"):
    ka = g(d, "Knowledge Area")
    if not ka:
        continue
    cells = {k: v for k, v in d.items() if k not in ("Knowledge Area", "Scope Note")}
    ka_matrix.append({"knowledgeArea": ka, "scopeNote": g(d, "Scope Note"), "cells": cells})
model["kaMatrix"] = ka_matrix

# --- Governance chain ---
model["rulesGov"] = [{
    "id": g(d, "Rule ID"), "name": g(d, "Name"), "scope": g(d, "Rule Scope"),
    "statement": g(d, "Rule Statement"), "appliesTo": g(d, "Applies To"), "requirement": g(d, "Requirement"),
} for d in rows("Rules") if g(d, "Rule ID")]

model["controls"] = [{
    "id": g(d, "Control ID"), "name": g(d, "Name"), "controlType": g(d, "Control Type"),
    "appliesTo": g(d, "Applies To"), "objective": g(d, "Control Objective / Evaluation"),
    "outcome": g(d, "Outcome"),
} for d in rows("Controls") if g(d, "Control ID")]

model["roles"] = [{
    "id": g(d, "Role ID"), "name": g(d, "Role Name"), "accountability": g(d, "Accountability Type"),
    "responsibility": g(d, "Responsibility"), "appliesTo": g(d, "Applies To"),
} for d in rows("Roles") if g(d, "Role ID")]

model["decisionRights"] = [{
    "id": g(d, "Decision Right ID"), "name": g(d, "Name"), "holder": g(d, "Holder Role ID"),
    "appliesTo": g(d, "Applies To"), "definition": g(d, "Decision Definition"), "requirement": g(d, "Requirement"),
} for d in rows("Decision Rights") if g(d, "Decision Right ID")]

model["exceptions"] = [{
    "id": g(d, "Exception ID"), "name": g(d, "Name"), "transition": g(d, "Applies To Transition"),
    "basis": g(d, "Trigger / Basis"), "authority": g(d, "Approving Authority"),
    "conditions": g(d, "Required Conditions"), "statusValues": g(d, "Lifecycle Status Values"),
} for d in rows("Exceptions") if g(d, "Exception ID")]

model["artefacts"] = [{
    "id": g(d, "Artefact ID"), "name": g(d, "Name"), "artefactType": g(d, "Artefact Type"),
    "producedIn": g(d, "Produced / Modified In"), "producingActivity": g(d, "Producing Activity"),
    "evidenceUse": g(d, "Evidence Use"),
} for d in rows("Artefacts") if g(d, "Artefact ID")]

model["evidence"] = [{
    "id": g(d, "Evidence ID"), "name": g(d, "Name"), "evidenceType": g(d, "Evidence Type"),
    "relatesTo": g(d, "Relates To"), "description": g(d, "Evidence Description"), "requirement": g(d, "Requirement"),
} for d in rows("Evidence") if g(d, "Evidence ID")]

model["compositeStates"] = [{
    "id": g(d, "Composite ID"), "name": g(d, "Name"), "globalState": g(d, "Global State ID"),
    "subState": g(d, "Global Sub-State ID"), "kaDimensions": g(d, "Included KA State Dimensions"),
    "context": g(d, "Context / Point in Time"), "vectorType": g(d, "Vector Type"),
} for d in rows("Composite States") if g(d, "Composite ID")]

model["relationships"] = [{
    "id": g(d, "Relationship ID"), "source": g(d, "Source Element ID"),
    "type": g(d, "Relationship Type"), "target": g(d, "Target Element ID"),
    "cardinality": g(d, "Cardinality"), "optionality": g(d, "Optionality"), "rationale": g(d, "Rationale / Notes"),
} for d in rows("Relationships") if g(d, "Relationship ID")]

# --- counts for sanity ---
counts = {k: len(v) for k, v in model.items() if isinstance(v, list)}
model["meta"]["counts"] = counts

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(model, f, ensure_ascii=False, indent=2)

print("Wrote", OUT)
print(json.dumps(counts, indent=2))
