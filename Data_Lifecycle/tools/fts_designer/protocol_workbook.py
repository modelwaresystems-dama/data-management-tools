#!/usr/bin/env python3
"""protocol_workbook.py: export a schema v0.3 model (three-layer Global Data Asset Architecture) to the
review workbook: README, Architecture, Guardrails, Lifecycle Phases, Lifecycle Activities, Regions, States,
Transitions, Events, Transition Guards, Cross-Region Constraints, State Vectors, Entry Conditions,
Exit Conditions, Invariants, GRCA Services, Service Families, Controls, Decision Rights, Permission Records,
Exceptions, Evidence, Conformance Clauses, Recommendations, Editorial Decisions, Naming Standard,
Name QA, Element Register, Relationships, Source Register, Derivation Trace, QA Findings, Legend.

Usage: python protocol_workbook.py model.fts.json out.xlsx
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font
from fts_workbook import sheet, g, TITLE_FONT, WRAP

def export(m, out):
    meta = m["meta"]; src = meta.get("source", "")
    S = lambda d: g(d, "status", meta.get("status", ""))
    O = lambda d: g(d, "origin", src)
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("README")
    lines = [
        (f"{meta.get('name','')} — {meta.get('modelId','')}", TITLE_FONT),
        (f"Model version {meta.get('version','')}  |  Schema {meta.get('schemaVersion','')}  |  Built {meta.get('buildStamp','')}  |  Status: {meta.get('status','')}", None),
        (f"Subject: {meta.get('subjectType','')}", None), (f"Sources: {src}", None), ("", None),
        ("What this workbook is", Font(bold=True)),
        (meta.get("note", ""), None), ("", None),
        ("The three layers", Font(bold=True)),
        *[(f"Layer {l['id'][-1]} {l['name']}: {l['question']} Constructs: {l['constructs']} Boundary: {l['boundary']}", None) for l in meta.get("layers", [])],
        ("", None), ("How to review", Font(bold=True)),
        (("1. Regions are the Knowledge Area's FTSs: one state machine per managed element (see the Regions sheet), running concurrently; the KA's condition at a time is one state per region (the State Vector)." if meta.get("level") == "knowledge-area" else "1. Regions and States are the Global FTS: four machines that run concurrently; the Data Asset's condition at a time is one state per region (the State Vector)."), None),
        ("2. Transitions, Events, Guards and Cross-Region Constraints are the transition contract: Eligible = Event and Active(source) and Exit and Guard and Entry and Invariants; Permitted = Eligible and risk acceptability and control satisfaction and assurance satisfaction and authorization.", None),
        ("3. Lifecycle Phases and Activities are Layer 1: work, never states. Permission Records say what an activity may do in a given state or vector.", None),
        ("4. GRCA Services, Controls, Decision Rights, Exceptions and Evidence are Layer 3.", None),
        ("5. QA Findings lists what the builder and the validation engine noticed. Rows whose Source names the encrypted workbook are stand-ins until that workbook can be read.", None),
        ("6. Keep IDs stable when editing. Every row is Proposed / illustrative until the Editorial Board says otherwise.", None),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t); c.alignment = WRAP
        if f: c.font = f
    ws.column_dimensions["A"].width = 130

    if meta.get("layers"): sheet(wb, "Architecture", ["Layer", "Name", "Primary question", "Primary constructs", "Boundary"], [[l["id"], l["name"], l["question"], l["constructs"], l["boundary"]] for l in meta.get("layers", [])], [8, 40, 50, 50, 50])
    cc = m.get("contextCapture") or {}
    if cc:
        rows = [["Definition", "", cc.get("definition", ""), "", ""], ["Ensures", "", cc.get("ensures", ""), "", ""]]
        for gl in cc.get("goals", []): rows.append(["Goal", "", gl, "", ""])
        for d in cc.get("businessDrivers", []): rows.append(["Business driver", "", d, "", ""])
        for i, x in enumerate(cc.get("inputs", []), 1): rows.append(["Input", f"IN-{i}", x, "", ""])
        for pr in cc.get("processes", []):
            rows.append(["Process", pr.get("id"), pr.get("name"), pr.get("phase", ""), pr.get("note", "")])
            for sa in pr.get("subActivities", []): rows.append(["Sub-activity", sa.split(" ")[0], sa[len(sa.split(" ")[0]):].strip(), pr.get("phase", ""), ""])
        for i, x in enumerate(cc.get("deliverables", []), 1): rows.append(["Deliverable", f"DEL-{i}", x, "", ""])
        for sec in ["suppliers", "participants", "consumers", "techniques", "tools", "metrics"]:
            for i, x in enumerate(cc.get(sec, []), 1): rows.append([sec[:-1].capitalize(), f"{sec[:3].upper()}-{i}", x, "", ""])
        sheet(wb, "Context Capture", ["Section", "Item ID", "Item", "Phase", "Note"], rows, [16, 10, 80, 8, 50], note=f"Structured capture of the {cc.get('knowledgeArea', '')} context diagram (image pages). {cc.get('phaseTags', '')}")
    if m.get("contributions"):
        sheet(wb, "Global Contributions", ["Contribution ID", "Global transition", "Kind", "Required KA states", "Predicate", "Simulator expression", "Requirement", "Note", "Status", "Source"],
              [[c["id"], c["globalTransition"], c["kind"], json.dumps(c.get("requiredStates", {})), c["predicate"], c.get("expression") or "", c.get("requirement", ""), c.get("note", ""), S(c), O(c)] for c in m["contributions"]], [14, 14, 12, 40, 70, 40, 14, 50, 20, 40],
              note="How this Knowledge Area reaches the Global Data Asset Protocol: guards, decision rights, services and events on named Global transitions (GA-005, GA-009). The viewer federates the guard rows onto the Global transitions; the simulator derives the fact values from this model's State Vector when both models are loaded.")
    if m.get("kaCouplings"):
        sheet(wb, "KA Couplings", ["Coupling ID", "Target model", "Target transition", "Event", "Expression", "Predicate", "Note", "Status", "Source"],
              [[k["id"], k["targetModel"], k["targetTransition"], k.get("event", ""), k.get("expression", ""), k["predicate"], k.get("note", ""), S(k), O(k)] for k in m["kaCouplings"]], [14, 12, 16, 12, 34, 70, 50, 20, 40],
              note="How this Knowledge Area couples to another Knowledge Area's FTS (events emitted into it, or its facts cited by a transition here). Loose coupling: never a shared state, never a Global transition (those are on Global Contributions).")
    if m.get("regulatoryFacts"):
        sheet(wb, "Regulatory Facts", ["Fact ID", "Regulation", "Topic", "Obligation", "Model reference", "Verification", "Status", "Source"],
              [[r["id"], r["regulation"], r["topic"], r["obligation"], r["modelReference"], r.get("verification", ""), S(r), O(r)] for r in m["regulatoryFacts"]], [14, 18, 30, 80, 40, 12, 20, 40],
              note="Regulatory obligations the regulation-neutral states, events and time triggers cite. Periods marked verify are to be checked against the current text of each Act before use.")
    if m.get("roles"):
        sheet(wb, "Roles", ["Role ID", "Role", "Category", "Responsibility", "Status", "Source"], [[r["id"], r["name"], g(r, "accountability"), g(r, "responsibility"), S(r), O(r)] for r in m["roles"]], [12, 30, 12, 60, 20, 40])
    if m.get("artefacts"):
        sheet(wb, "Artefacts", ["Artefact ID", "Artefact", "Produced in", "Producing activity", "Evidence use", "Status", "Source"], [[a["id"], a["name"], g(a, "producedIn"), g(a, "producingActivity"), g(a, "evidenceUse"), S(a), O(a)] for a in m["artefacts"]], [12, 40, 14, 40, 50, 20, 40])
    if m.get("rules"): sheet(wb, "Guardrails", ["ID", "Guardrail", "Statement", "Status", "Authority"], [[r["id"], r["shortName"], r["statement"], r["status"], r["authority"]] for r in m["rules"]], [8, 30, 90, 22, 40])
    ph = m.get("lifecycle", {}).get("phases", [])
    if ph: sheet(wb, "Lifecycle Phases", ["Phase ID", "Phase", "Purpose", "Execution trait / concurrency", "Entry context", "Activities", "Name QA", "Status", "Source"],
          [[p["id"], p["name"], p["purpose"], p["nonLinearity"], g(p, "entryContext"), "; ".join(a["id"] for a in m["activities"] if p["id"] in a.get("lifecyclePhases", [])), g(g(p, "nameQA", {}), "status"), S(p), O(p)] for p in ph], [10, 26, 60, 34, 40, 40, 10, 20, 40],
          note=(m.get("lifecycle", {}).get("mappingRule") or ""))
    if m.get("lifecycle", {}).get("effectClasses"): sheet(wb, "Effect Classes", ["Effect classification", "Definition", "Examples", "Source"], [[e["effectClass"], e["definition"], e["examples"], e.get("origin", "")] for e in m.get("lifecycle", {}).get("effectClasses", [])], [24, 90, 60, 50])
    sheet(wb, "Activities" if not ph else "Lifecycle Activities", ["Activity ID", "Activity", "Effect classification", "Lifecycle phases", "Related transitions", "GRCA services", "Permission records", "Name QA", "Notes", "Status", "Source"],
          [[a["id"], a["name"], a["effectClass"], a.get("lifecyclePhases", []), a.get("relatedTransitions", []), a.get("services", []), "; ".join(p["id"] for p in m.get("permissionRecords", []) if p["activity"] == a["id"]), g(g(a, "nameQA", {}), "status"), g(a, "notes"), S(a), O(a)] for a in m["activities"]], [10, 26, 34, 24, 30, 34, 20, 10, 50, 20, 40])
    ka_regions = any(r.get("managedElement") for r in m.get("regions", []))
    if ka_regions:
        sheet(wb, "Regions", ["Region ID", "Region (FTS)", "Code", "Managed element", "Element kind", "Instance scope", "Conditions that matter", "Question", "Initial state", "Terminal states", "Region invariant / rule", "States", "Transitions", "Contributes to the Global regions", "Issue sources", "Status", "Source"],
              [[r["id"], r["name"], r["code"], r.get("managedElement", ""), r.get("elementKind", ""), r.get("instanceScope", ""), r.get("conditionsThatMatter", ""), r["question"], r["initialState"], "; ".join(s["id"] for s in m["subStates"] if s.get("region") == r["id"] and s.get("terminal")), r["regionInvariant"], sum(1 for s in m["subStates"] if s.get("region") == r["id"]), sum(1 for t in m["transitions"] if t.get("region") == r["id"] and t["level"] != "Initial"), r.get("contributesTo", ""), "; ".join(r.get("issueSources", [])), S(r), O(r)] for r in m.get("regions", [])], [10, 24, 6, 44, 14, 30, 36, 44, 12, 14, 44, 8, 10, 60, 40, 20, 40],
              note=(m.get("meta", {}).get("regionModel") or "One FTS per managed element: each region is its own state machine over the element it names."))
    else:
        sheet(wb, "Regions", ["Region ID", "Region", "Code", "Question", "Initial state", "Terminal states", "Region invariant / rule", "States", "Transitions", "Status", "Source"],
          [[r["id"], r["name"], r["code"], r["question"], r["initialState"], "; ".join(s["id"] for s in m["subStates"] if s.get("region") == r["id"] and s.get("terminal")), r["regionInvariant"], sum(1 for s in m["subStates"] if s.get("region") == r["id"]), sum(1 for t in m["transitions"] if t.get("region") == r["id"] and t["level"] != "Initial"), S(r), O(r)] for r in m.get("regions", [])], [10, 24, 6, 50, 12, 20, 50, 8, 10, 20, 40])
    sheet(wb, "States", ["State ID", "Region", "State", "Initial?", "Terminal?", "Definition", "Invariant", "Entry condition", "Exit condition", "Semantic class", "Name QA", "Name QA rationale", "Status", "Source"],
          [[s["id"], s.get("region"), s["name"], s.get("initial", False), s.get("terminal", False), s["definition"],
            "; ".join(i["predicate"] for i in m["invariants"] if i["appliesTo"] == s["id"]), "; ".join(i["predicate"] for i in m["entryConditions"] if i["appliesTo"] == s["id"]), "; ".join(i["predicate"] for i in m["exitConditions"] if i["appliesTo"] == s["id"]),
            g(s, "semanticClass"), g(g(s, "nameQA", {}), "status"), g(g(s, "nameQA", {}), "rationale"), S(s), O(s)] for s in m["subStates"]], [12, 10, 24, 8, 8, 60, 50, 50, 40, 20, 10, 50, 20, 40])
    sheet(wb, "Transitions", ["Transition ID", "Transition", "Region", "Source", "Target", "Event", "Kind", "Guard summary", "Cross-region constraints", "Decision right", "GRCA services", "Reversibility", "Evidence", "Name QA", "Status", "Source"],
          [[t["id"], t["name"], t.get("region"), t["source"], t["target"], t.get("event"), t.get("transitionKind"), t.get("guardSummary"), t.get("crossRegionConstraints", []), t.get("decisionRight") or "None", t.get("services", []), t.get("reversibility"),
            "; ".join(e["id"] for e in m.get("evidence", []) if e.get("relatesTo") == t["id"]), g(g(t, "nameQA", {}), "status"), S(t), O(t)] for t in m["transitions"]], [12, 30, 10, 12, 12, 10, 9, 60, 22, 12, 34, 24, 14, 10, 20, 40])
    sheet(wb, "Events", ["Event ID", "Event", "Event type", "Meaning", "Notes", "Status", "Source"], [[e["id"], e["name"], g(e, "eventType"), g(e, "meaning"), g(e, "notes"), S(e), O(e)] for e in m["events"]], [10, 34, 20, 50, 50, 20, 40])
    sheet(wb, "Transition Guards", ["Guard ID", "Guard", "Transition", "Predicate", "Scope", "Requirement", "Expression (simulator)", "Constraint", "Status", "Source"],
          [[x["id"], x["name"], x["transition"], x["predicate"], g(x, "scope"), g(x, "requirement"), g(x, "expression") or "", g(x, "constraint"), S(x), O(x)] for x in m["guards"]], [18, 34, 12, 60, 12, 24, 50, 10, 20, 40])
    sheet(wb, "Cross-Region Constraints", ["Constraint ID", "Constraint", "Transitions", "Requirement", "Expression (simulator)", "Expression note", "Status", "Source"],
          [[x["id"], x["constraint"], x["transitions"], x["requirement"], x.get("expression"), g(x, "expressionNote"), S(x), O(x)] for x in m.get("crossRegionConstraints", [])], [12, 70, 22, 22, 60, 50, 20, 40])
    sheet(wb, "State Vectors", ["Vector ID", "Configuration", "Existence", "Assurance", "Availability", "Custody / Preservation", "Legality / use note", "Status", "Source"],
          [[v["id"], v["name"], v["vector"].get("REG-EX"), v["vector"].get("REG-AS"), v["vector"].get("REG-AV"), v["vector"].get("REG-CP"), v["legality"], S(v), O(v)] for v in m.get("stateVectors", [])], [10, 30, 14, 14, 14, 14, 60, 20, 40])
    sheet(wb, "Entry Conditions", ["Condition ID", "Applies to", "Predicate", "Requirement", "Notes", "Status", "Source"], [[c["id"], c["appliesTo"], c["predicate"], g(c, "requirement"), g(c, "notes"), S(c), O(c)] for c in m["entryConditions"]], [10, 12, 80, 12, 50, 20, 40])
    sheet(wb, "Exit Conditions", ["Condition ID", "Applies to", "Predicate", "Requirement", "Status", "Source"], [[c["id"], c["appliesTo"], c["predicate"], g(c, "requirement"), S(c), O(c)] for c in m["exitConditions"]], [10, 12, 80, 12, 20, 40])
    sheet(wb, "Invariants", ["Invariant ID", "Applies to", "Predicate", "Severity", "Status", "Source"], [[c["id"], c["appliesTo"], c["predicate"], g(c, "severity"), S(c), O(c)] for c in m["invariants"]], [10, 12, 80, 10, 20, 40])
    sheet(wb, "GRCA Services", ["Service ID", "Family", "Service", "Trigger", "Output", "Used by transitions", "Used by activities", "Status", "Source"],
          [[s["id"], s["family"], s["name"], s["trigger"], s["output"], s.get("usedByTransitions", []), s.get("usedByActivities", []), S(s), O(s)] for s in m.get("services", [])], [12, 12, 34, 40, 40, 40, 30, 20, 40])
    sheet(wb, "Service Families", ["Family", "Fundamental question", "Typical outputs"], [[f["family"], f["question"], f["outputs"]] for f in m.get("serviceFamilies", [])], [14, 60, 70])
    sheet(wb, "Controls", ["Control ID", "Control", "Type", "Applies to", "Trigger / objective", "Result", "Service", "Policy domain", "Control number", "Theme", "Implementing procedure (template)", "Evidence artefact (template)", "Status", "Source"],
          [[c["id"], c["name"], g(c, "controlType"), g(c, "appliesTo"), g(c, "objective"), g(c, "outcome"), g(c, "service"), g(c, "policyDomain"), g(c, "controlNumber"), g(c, "theme"), g(c, "procedure"), g(c, "evidenceArtefact"), S(c), O(c)] for c in m["controls"]],
          [12, 40, 22, 30, 40, 30, 12, 14, 10, 30, 12, 34, 20, 40],
          note="Mechanism: a control service the model uses. Policy control: a control of the Knowledge Area policy in the FutureState workbooks, named by domain and number (State Contracts register, 25 Sep 2026); the twin resolves each organisation's own wording, procedure and evidence artefact.")
    # State Contracts register card 7 option a (Howard, 25 Sep 2026): the drafted mapping, one row per transition, for review
    cid = {c["id"]: c for c in m["controls"]}
    tcm = []
    for t in m["transitions"]:
        if t.get("level") == "Initial" or t.get("source") == "[Initial]": continue
        pcs = t.get("policyControls")
        if pcs is None and not t.get("policyControlGap"): continue
        pcs = pcs or []
        tcm.append([t["id"], t["name"], g(t, "region"), "; ".join(x["control"] for x in pcs), "; ".join(cid.get(x["control"], {}).get("name", "") for x in pcs),
                    "; ".join(sorted({str(cid.get(x["control"], {}).get("procedure") or "") for x in pcs} - {""})), "; ".join(cid.get(x["control"], {}).get("evidenceArtefact") or "" for x in pcs),
                    (pcs[0]["why"] if pcs else t.get("policyControlGap", "")), "Proposed (drafted 25 Sep 2026)" if pcs else "Gap: no control fits", ""])
    if tcm:
        sheet(wb, "Transition Control Map", ["Transition ID", "Transition", "Region", "Policy controls", "Control wording (template)", "Implementing procedures (template)", "Evidence artefacts (template)", "Why", "Status", "Your review"],
              tcm, [12, 30, 16, 26, 60, 14, 40, 60, 22, 30],
              note="Drafted for Howard's review (State Contracts register card 7 option a, 25 Sep 2026). Change a mapping in the Knowledge Area spec (POLICY_CONTROLS) and rebuild; the twin checks that the procedure implementing each control is in force (card 9).")
    sheet(wb, "Decision Rights", ["Decision Right ID", "Decision right", "Holder role", "Applies to transitions", "Definition", "Service", "Notes", "Status", "Source"], [[d["id"], d["name"], g(d, "holder"), g(d, "appliesTo"), g(d, "definition"), g(d, "service"), g(d, "notes"), S(d), O(d)] for d in m["decisionRights"]], [12, 40, 14, 30, 60, 12, 50, 20, 40])
    sheet(wb, "Permission Records", ["Permission ID", "Activity", "Context (state / vector)", "Outcome", "Guard", "Authority", "Services", "Evidence", "Effect", "Basis", "Status", "Source"],
          [[p["id"], p["activity"], p["context"], p["outcome"], p["guard"], p["authority"], p.get("services", []), p.get("evidence"), p.get("effect"), p.get("basis"), S(p), O(p)] for p in m.get("permissionRecords", [])], [10, 10, 16, 14, 60, 30, 34, 30, 26, 30, 20, 40],
          note="The report's worked examples and the definitions of the Availability states. The full permission matrix is in the validated workbook, which could not be read (encrypted).")
    sheet(wb, "Exceptions", ["Exception ID", "Exception", "Transition", "Basis", "Authority", "Conditions", "Status values", "Status", "Source"], [[x["id"], x["name"], g(x, "transition"), g(x, "basis"), g(x, "authority"), g(x, "conditions"), g(x, "statusValues"), S(x), O(x)] for x in m["exceptions"]], [10, 36, 12, 40, 30, 70, 30, 20, 40])
    sheet(wb, "Evidence", ["Evidence ID", "Evidence", "Type", "Relates to", "Description", "Requirement", "Status", "Source"], [[e["id"], e["name"], g(e, "evidenceType"), ", ".join(e.get("transitions") or []) or g(e, "relatesTo"), g(e, "description"), g(e, "requirement"), S(e), O(e)] for e in m["evidence"]], [10, 36, 22, 12, 70, 12, 20, 40])
    if m.get("conformance"): sheet(wb, "Conformance Clauses", ["Clause ID", "Title", "Candidate normative text (SHALL is proposed, not approved)"], [[c["id"], c["title"], c["clause"]] for c in m.get("conformance", [])], [10, 30, 110])
    if m.get("recommendations"): sheet(wb, "Recommendations", ["ID", "Recommendation", "Proposed decision", "Suggested disposition"], [[r["id"], r["title"], r["proposedDecision"], r["suggestedDisposition"]] for r in m.get("recommendations", [])], [8, 40, 80, 26])
    if m.get("editorialDecisions"): sheet(wb, "Editorial Decisions", ["Decision ID", "Question for the Board", "Recommended decision", "If deferred"], [[e["id"], e["question"], e["recommended"], e["ifDeferred"]] for e in m.get("editorialDecisions", [])], [10, 50, 60, 50])
    if m.get("namingStandard"): sheet(wb, "Naming Standard", ["Element type", "Naming pattern", "Example", "Source"], [[n["elementType"], n["pattern"], n["example"], n["source"]] for n in m.get("namingStandard", [])], [20, 70, 40, 60])
    qa = [[s["id"], "State", s["name"], g(g(s, "nameQA", {}), "status"), g(g(s, "nameQA", {}), "rationale")] for s in m["subStates"]] + \
         [[t["id"], "Transition", t["name"], g(g(t, "nameQA", {}), "status"), g(g(t, "nameQA", {}), "rationale")] for t in m["transitions"] if t["level"] != "Initial"] + \
         [[a["id"], "Activity", a["name"], g(g(a, "nameQA", {}), "status"), g(g(a, "nameQA", {}), "rationale")] for a in m["activities"]] + \
         [[p["id"], "Lifecycle Phase", p["name"], g(g(p, "nameQA", {}), "status"), g(g(p, "nameQA", {}), "rationale")] for p in ph]
    sheet(wb, "Name QA", ["Element ID", "Element type", "Name", "QA result", "Rationale"], qa, [12, 16, 34, 10, 90], note="Semantic test (durable condition for states; behaviour for activities and transitions); grammar is the heuristic, per the naming standard.")
    reg = []
    for p in ph: reg.append([p["id"], "Lifecycle Phase", p["name"], "", p["purpose"], "L1", S(p), O(p)])
    for a in m["activities"]: reg.append([a["id"], "Activity", a["name"], "; ".join(a.get("lifecyclePhases", [])), a["effectClass"], "L1", S(a), O(a)])
    for r in m.get("regions", []): reg.append([r["id"], "Region", r["name"], "", r["question"], "L2", S(r), O(r)])
    for s in m["subStates"]: reg.append([s["id"], "State", s["name"], s.get("region"), s["definition"], "L2", S(s), O(s)])
    for k, typ, pk, dk, layer in [("transitions", "Transition", "region", "guardSummary", "L2"), ("events", "Event", "", "meaning", "L2"), ("guards", "Guard", "transition", "predicate", "L2"), ("crossRegionConstraints", "Cross-Region Constraint", "", "constraint", "L2"), ("stateVectors", "State Vector", "", "legality", "L2"), ("invariants", "Invariant", "appliesTo", "predicate", "L2"), ("entryConditions", "Entry Condition", "appliesTo", "predicate", "L2"), ("exitConditions", "Exit Condition", "appliesTo", "predicate", "L2"), ("services", "GRCA Service", "family", "output", "L3"), ("controls", "Control", "appliesTo", "objective", "L3"), ("decisionRights", "Decision Right", "appliesTo", "definition", "L3"), ("permissionRecords", "Permission Record", "context", "guard", "L3"), ("exceptions", "Exception", "transition", "basis", "L3"), ("evidence", "Evidence", "relatesTo", "description", "L3")]:
        for r in m.get(k, []): reg.append([r["id"], typ, g(r, "name", g(r, "constraint", g(r, "predicate", ""))), g(r, pk) if pk else "", g(r, dk), layer, S(r), O(r)])
    sheet(wb, "Element Register", ["Element ID", "Element type", "Name", "Parent / applies to", "Definition / purpose", "Layer", "Status", "Source"], reg, [14, 22, 34, 18, 70, 6, 20, 40])
    sheet(wb, "Relationships", ["Relationship ID", "Source", "Relationship", "Target", "Cardinality", "Optionality", "Rationale", "Status"], [[r["id"], r["source"], r["type"], r["target"], g(r, "cardinality"), g(r, "optionality"), g(r, "rationale"), S(r)] for r in m["relationships"]], [12, 16, 30, 16, 10, 12, 40, 20])
    sheet(wb, "Source Register", ["Source ID", "Source", "Type", "Location", "Use", "Limitations"], [[s["id"], s["source"], s["type"], s["location"], s["use"], s["limitations"]] for s in m.get("sources", [])], [10, 60, 20, 30, 60, 60])
    sheet(wb, "Derivation Trace", ["Element ID", "Element type", "Name / predicate", "Derived from"], [[t["elementId"], t["elementType"], t["name"], t["derivedFrom"]] for t in m.get("derivationTrace", [])], [12, 20, 50, 50])
    rv = m.get("reviewItems", [])
    if rv: sheet(wb, "Review Items", ["ID", "When", "Model version", "Element type", "Element ID", "Element", "Verdict", "Note", "State vector at the time", "Known element"], [[r.get("id"), r.get("at"), r.get("modelVersion"), r.get("elementType"), r.get("elementId"), r.get("elementName"), r.get("verdict"), r.get("note"), json.dumps((r.get("context") or {}).get("vector")) if (r.get("context") or {}).get("vector") else "", r.get("known", True)] for r in rv], [16, 20, 10, 18, 14, 34, 10, 70, 40, 8], note="Howard's feedback captured in the viewer (reviews/*_feedback.json) and attached by fts_feedback.py.")
    sheet(wb, "QA Findings", ["Severity", "Rule", "Element", "Finding"], [[f["severity"], f["rule"], f["element"], f["finding"]] for f in m.get("qaFindings", [])], [10, 10, 20, 100], note="Builder findings. The validation engine's findings are in the simulation report workbook.")
    leg = [["T" + str(i + 1), t["id"], t["name"]] for i, t in enumerate(m["transitions"])] + [["E" + str(i + 1), e["id"], e["name"]] for i, e in enumerate(m["events"])] + [["G" + str(i + 1), x["id"], x["name"]] for i, x in enumerate(m["guards"])] + [["A" + str(i + 1), a["id"], a["name"]] for i, a in enumerate(m["activities"])]
    sheet(wb, "Legend", ["Code", "Element ID", "Name"], leg, [8, 18, 60], note="Short codes used by the simplified diagram.")
    wb.save(out); return out

if __name__ == "__main__":
    m = json.load(open(sys.argv[1], encoding="utf-8")); print("Wrote", export(m, sys.argv[2]))
