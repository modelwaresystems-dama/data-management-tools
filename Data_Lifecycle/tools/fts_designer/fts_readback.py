#!/usr/bin/env python3
"""
fts_readback.py  -  reads Howard's edits from a review workbook back into an overrides file that the spec build
applies, so the workbook is the review surface and the spec-as-code stays the source.

Reads these sheets and columns (others are ignored):
  Decision Rights : Decision right, Holder role (role ID or role name from the Roles sheet), Notes
  States          : State, Definition, Invariant
  Transitions     : Transition, Event, Guard summary, Decision right
  Events          : Event, Event type
  Roles           : Role, Responsibility

Usage:
  python fts_readback.py models/<name>.fts.json models/<name>_workbook.xlsx [--out spec/<name>_overrides.json]
Then rebuild:  python spec/<ka>_spec.py models/ --overrides spec/<name>_overrides.json

The overrides file lists only cells that differ from the model. Each override records the old and new value;
ka_build applies the new value, marks the element's trace "workbook:<sheet>" and clears a REVIEW note when the
holder was changed. Run it again after the next review; overrides are cumulative per element and field.
"""
import argparse, datetime, json, os, sys
try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

SHEETS = {
    "Decision Rights": {"key": "Decision Right ID", "collection": "decisionRights", "fields": {"Decision right": "name", "Holder role": "holder", "Notes": "notes"}},
    "States": {"key": "State ID", "collection": "subStates", "fields": {"State": "name", "Definition": "definition", "Invariant": "invariant"}},
    "Transitions": {"key": "Transition ID", "collection": "transitions", "fields": {"Transition": "name", "Event": "event", "Guard summary": "guardSummary", "Decision right": "decisionRight"}},
    "Events": {"key": "Event ID", "collection": "events", "fields": {"Event": "name", "Event type": "eventType"}},
    "Roles": {"key": "Role ID", "collection": "roles", "fields": {"Role": "name", "Responsibility": "responsibility"}},
}

def norm(v):
    if v is None: return ""
    s = str(v).strip()
    return "" if s in ("None", "") else s

def read_sheet(wb, name):
    if name not in wb.sheetnames: return None, []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    # find the header row (first row whose first cell is the key column)
    hdr_i = next((i for i, r in enumerate(rows) if r and norm(r[0]) == SHEETS[name]["key"]), None)
    if hdr_i is None: return None, []
    hdr = [norm(c) for c in rows[hdr_i]]
    return hdr, [dict(zip(hdr, r)) for r in rows[hdr_i + 1:] if r and norm(r[0])]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("model"); ap.add_argument("workbook"); ap.add_argument("--out")
    a = ap.parse_args()
    m = json.load(open(a.model, encoding="utf-8"))
    wb = openpyxl.load_workbook(a.workbook, data_only=True)
    out_path = a.out or os.path.splitext(a.model)[0].replace(".fts", "") + "_overrides.json"
    prior = json.load(open(out_path, encoding="utf-8")) if os.path.exists(out_path) else {"overrides": []}
    existing = {(o["collection"], o["id"], o["field"]): o for o in prior.get("overrides", [])}
    roles_by_name = {r["name"].strip().lower(): r["id"] for r in m.get("roles", [])}
    inv_by_state = {i["appliesTo"]: i for i in m.get("invariants", [])}
    by_id = {c: {x["id"]: x for x in m.get(c, [])} for c in {s["collection"] for s in SHEETS.values()}}
    changes = []
    for sheet, spec in SHEETS.items():
        hdr, rows = read_sheet(wb, sheet)
        if hdr is None: continue
        for r in rows:
            eid = norm(r.get(spec["key"]))
            cur = by_id[spec["collection"]].get(eid)
            if not cur: continue
            for col, field in spec["fields"].items():
                if col not in hdr: continue
                new = norm(r.get(col))
                if field == "invariant":
                    old = norm(inv_by_state.get(eid, {}).get("predicate"))
                elif field == "holder":
                    old = norm(cur.get("holder"))
                    if new and new not in {x["id"] for x in m.get("roles", [])}:
                        new = roles_by_name.get(new.lower(), new)   # accept a role name
                else:
                    old = norm(cur.get(field))
                if field == "decisionRight" and new in ("None", ""): new = ""
                if new == old: continue
                if not new and field in ("name", "definition", "event"): continue   # blank never erases a required value
                o = {"collection": spec["collection"], "id": eid, "field": field, "old": old, "new": new, "sheet": sheet, "at": datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=2))).strftime("%Y-%m-%dT%H:%M SAST")}
                existing[(spec["collection"], eid, field)] = o; changes.append(o)
    doc = {"model": m["meta"].get("modelId"), "modelVersion": m["meta"].get("version"), "workbook": os.path.basename(a.workbook), "overrides": list(existing.values())}
    json.dump(doc, open(out_path, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print(f"{len(changes)} new change(s), {len(doc['overrides'])} override(s) in {out_path}")
    for o in changes: print(f"  {o['sheet']}: {o['id']}.{o['field']}: {o['old']!r} -> {o['new']!r}")

if __name__ == "__main__":
    main()
