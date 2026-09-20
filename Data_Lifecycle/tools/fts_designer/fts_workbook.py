#!/usr/bin/env python3
"""fts_workbook.py: export a canonical .fts.json model to the 27-sheet review workbook
(same sheet order and columns as DMBOK_Plan_Design_State_Transition_Metadata_v2.xlsx),
plus Context Capture, Derivation Trace, QA Findings and Legend sheets.

Usage: python fts_workbook.py model.fts.json out.xlsx
"""
import json, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E5A"); HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=14); NOTE_FONT = Font(italic=True, size=9, color="555555")
THIN = Side(style="thin", color="D0D7DE"); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

def sheet(wb, title, headers, rows, widths=None, note=None):
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = NOTE_FONT; r0 = 2
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=j, value=h); c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = WRAP; c.border = BORDER
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (list, dict)): v = "; ".join(v) if isinstance(v, list) else json.dumps(v, ensure_ascii=False)
            if isinstance(v, bool): v = "Yes" if v else "No"
            c = ws.cell(row=i, column=j, value=v); c.alignment = WRAP; c.border = BORDER
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{max(r0+1, r0+len(rows))}"
    for j, h in enumerate(headers, 1):
        w = (widths[j-1] if widths and j-1 < len(widths) else None) or min(60, max(12, int(max([len(str(h))] + [len(str(r[j-1])) for r in rows if j-1 < len(r) and r[j-1] is not None][:200]) * 0.9)))
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws

def g(d, k, default=""):
    v = d.get(k, default); return default if v is None else v

def export(m, out):
    meta = m["meta"]; src = meta.get("source", "")
    S = lambda d: g(d, "status", meta.get("status", ""))
    wb = Workbook(); wb.remove(wb.active)

    # README
    ws = wb.create_sheet("README")
    lines = [
        (f"{meta.get('name','FTS model')} — {meta.get('modelId','')}", TITLE_FONT),
        (f"Level: {meta.get('level','')}" + (f"  |  Knowledge Area: {meta['knowledgeArea']}" if meta.get('knowledgeArea') else ""), None),
        (f"Version {meta.get('version','')}  |  Built {meta.get('buildStamp','')}  |  Status: {meta.get('status','')}", None),
        (f"Source: {src}", None), ("", None),
        ("Purpose", Font(bold=True)),
        ("Finite State Transition model of a Data Asset (or of a Knowledge Area's management condition). States and Sub-States are nouns describing conditions; Activities and Transitions are verbs. A transition is eligible when source exit conditions, transition guards and target entry conditions hold, and still needs the applicable Decision Right.", None),
        ("", None), ("How to review", Font(bold=True)),
        ("1. Read Important Retained Decisions: the rules this model is built to (normative for the model, not Editorial Board approvals).", None),
        ("2. Read QA Findings: every gap and doubtful name the compiler noticed. Nothing is blocked; findings are for your judgement.", None),
        ("3. Read Derivation Trace: where each derived row came from (context item, placement rule, or drafting).", None),
        ("4. Edit any sheet. Keep IDs stable. Re-import with tools/convert_fts.py to regenerate the model and diagrams.", None),
        ("5. Every row is Proposed / illustrative until the Board says otherwise (IRD-024).", None),
        ("", None), ("Sheets", Font(bold=True)),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t); c.alignment = WRAP
        if f: c.font = f
    ws.column_dimensions["A"].width = 120
    r = len(lines) + 1
    for name in ["Important Retained Decisions","State Name QA","Overview","Global States","Sub-States","Activities","Transitions","Events","Entry Conditions","Exit Conditions","Invariants","Transition Guards","KA Interactions","KA Matrix","Rules","Controls","Roles","Decision Rights","Exceptions","Artefacts","Evidence","Composite States","Element Register","Relationships","Data Dictionary","Source Register","Context Capture","Derivation Trace","QA Findings","Legend"]:
        ws.cell(row=r, column=1, value=name); r += 1

    sheet(wb, "Important Retained Decisions", ["Rule ID","Short Name","Rule Statement","Rationale","Applies To","Validation Test","Compliant Example","Non-Compliant Example","Status","Authority / Scope","Notes"],
          [[g(x,"id"),g(x,"shortName"),g(x,"statement"),g(x,"rationale"),g(x,"appliesTo"),g(x,"test"),g(x,"compliant"),g(x,"nonCompliant"),g(x,"status"),g(x,"authority"),g(x,"notes")] for x in m["rules"]],
          [10,24,50,40,26,40,30,30,18,30,30])

    states = m["globalStates"] + m["subStates"]
    def qa_result(s):
        st = g(g(s,"nameQA",{}),"status",""); return "PASS" if st.lower().startswith("pass") else "REVIEW" if st.lower().startswith("review") else st.upper()
    sheet(wb, "State Name QA", ["Element ID","Element Type","Name","Parent ID","State Test","Noun / Condition Test","Verb / Instruction Risk","QA Result","Rule IDs","Required Action","Notes"],
          [[s["id"], "Global State" if s["kind"]=="global" else "Sub-State", s["name"], g(s,"parent"),
            g(g(s,"nameQA",{}),"rationale",""), "Pass" if qa_result(s)=="PASS" else "Ambiguous noun/verb label", "Low" if qa_result(s)=="PASS" else "Medium", qa_result(s),
            "; ".join(g(s,"retainedRuleIds",[])), "None" if qa_result(s)=="PASS" else "Retain for this model; require formal state-test decision before any approval.", g(s,"notes")] for s in states],
          [14,14,24,14,50,24,14,10,28,40,30])

    counts = meta.get("counts", {})
    ov = [[f"{meta.get('name','')} — Overview"], [f"Model boundary", meta.get("note","")], ["Key caution","Important Retained Decisions are normative for this model, not DMBOK v3 approvals. All content is Proposed / illustrative."], ["Element Type","Count","Workbook Sheet"]]
    for k, lab, sh in [("globalStates","Global States","Global States"),("subStates","Sub-States","Sub-States"),("activities","Activities","Activities"),("transitions","Transitions","Transitions"),("events","Events","Events"),("entryConditions","Entry Conditions","Entry Conditions"),("exitConditions","Exit Conditions","Exit Conditions"),("invariants","Invariants","Invariants"),("guards","Transition Guards","Transition Guards"),("kaInteractions","KA Interactions","KA Interactions"),("rulesGov","Rules","Rules"),("controls","Controls","Controls"),("roles","Roles","Roles"),("decisionRights","Decision Rights","Decision Rights"),("exceptions","Exceptions","Exceptions"),("artefacts","Artefacts","Artefacts"),("evidence","Evidence","Evidence"),("compositeStates","Composite States","Composite States"),("relationships","Relationships","Relationships")]:
        ov.append([lab, counts.get(k, len(m.get(k,[]))), sh])
    ov.append(["QA findings", meta.get("qaCounts",{}).get("warnings",0), "warnings"]); ov.append(["", meta.get("qaCounts",{}).get("notes",0), "notes"])
    ws = wb.create_sheet("Overview")
    for i, row in enumerate(ov, 1):
        for j, v in enumerate(row, 1): ws.cell(row=i, column=j, value=v)
    ws["A1"].font = TITLE_FONT; ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 60; ws.column_dimensions["C"].width = 26

    sheet(wb, "Global States", ["State ID","Name","Definition / Purpose","State Type","Initial?","Terminal?","Status","Notes","Source","Semantic Classification","Name QA Status","Name QA Rationale","Retained Rule IDs"],
          [[s["id"],s["name"],g(s,"definition"),g(s,"stateType"),g(s,"initial",False),g(s,"terminal",False),S(s),g(s,"notes"),g(s,"origin",src),g(s,"semanticClass"),g(g(s,"nameQA",{}),"status"),g(g(s,"nameQA",{}),"rationale"),g(s,"retainedRuleIds",[])] for s in m["globalStates"]],
          [12,18,60,24,8,8,20,36,30,18,18,50,26])
    sheet(wb, "Sub-States", ["Sub-State ID","Name","Parent State ID","Definition / Purpose","Sequence Hint","Requirement","Readiness State?","Status","Notes","Source","Semantic Classification","Name QA Status","Name QA Rationale","Retained Rule IDs"],
          [[s["id"],s["name"],g(s,"parent"),g(s,"definition"),g(s,"sequence"),g(s,"requirement"),g(s,"readiness",False),S(s),g(s,"notes"),g(s,"origin",src),g(s,"semanticClass"),g(g(s,"nameQA",{}),"status"),g(g(s,"nameQA",{}),"rationale"),g(s,"retainedRuleIds",[])] for s in m["subStates"]],
          [14,22,12,60,8,26,10,20,36,30,18,12,50,26])
    sheet(wb, "Activities", ["Activity ID","Activity Name","Permitted In Element ID","Activity Type","Noun / Verb","Permissibility","Status","Notes","Source"],
          [[a["id"],a["name"],g(a,"permittedIn"),g(a,"activityType"),g(a,"nounVerb"),g(a,"permissibility"),S(a),g(a,"notes"),g(a,"origin",src)] for a in m["activities"]], [10,30,16,34,10,14,20,40,30])
    sheet(wb, "Transitions", ["Transition ID","Name","Level","Source Element ID","Target Element ID","Trigger Event ID","Transition Type","Optionality","Guard / Rationale Summary","Status","Source"],
          [[t["id"],t["name"],g(t,"level"),t["source"],t["target"],g(t,"event"),g(t,"transitionType"),g(t,"optionality"),g(t,"guardSummary"),S(t),g(t,"origin",src)] for t in m["transitions"]], [12,34,10,16,16,10,22,18,50,20,30])
    sheet(wb, "Events", ["Event ID","Name","Event Type","Definition / Trigger Meaning","Status","Source"],
          [[e["id"],e["name"],g(e,"eventType"),g(e,"meaning"),S(e),g(e,"origin",src)] for e in m["events"]], [10,40,20,60,20,30])
    sheet(wb, "Entry Conditions", ["Condition ID","Applies To Element ID","Predicate / Condition","Requirement","Condition Type","Status","Source"],
          [[c["id"],c["appliesTo"],c["predicate"],g(c,"requirement"),g(c,"conditionType"),S(c),g(c,"origin",src)] for c in m["entryConditions"]], [10,16,70,26,16,20,30])
    sheet(wb, "Exit Conditions", ["Condition ID","Applies To Element ID","Predicate / Condition","Requirement","Status","Source"],
          [[c["id"],c["appliesTo"],c["predicate"],g(c,"requirement"),S(c),g(c,"origin",src)] for c in m["exitConditions"]], [10,16,70,26,20,30])
    sheet(wb, "Invariants", ["Invariant ID","Applies To Element ID","Invariant Predicate","Severity","Status","Source"],
          [[c["id"],c["appliesTo"],c["predicate"],g(c,"severity"),S(c),g(c,"origin",src)] for c in m["invariants"]], [10,16,70,10,20,30])
    sheet(wb, "Transition Guards", ["Guard ID","Name","Transition ID","Predicate / Guard","Contributing Scope","Requirement","Status","Source"],
          [[x["id"],x["name"],x["transition"],x["predicate"],g(x,"scope"),g(x,"requirement"),S(x),g(x,"origin",src)] for x in m["guards"]], [12,28,14,60,22,22,20,30])
    sheet(wb, "KA Interactions", ["Interaction ID","Global Sub-State ID","Knowledge Area","Interaction Type","Applicability","Interaction Description","Status","Notes","Source"],
          [[k["id"],k["subState"],k["knowledgeArea"],g(k,"interactionType"),g(k,"applicability"),g(k,"description"),S(k),g(k,"notes"),g(k,"origin",src)] for k in m["kaInteractions"]], [10,16,30,20,28,60,20,40,30])
    cols = []
    for r in m["kaMatrix"]:
        for c in r.get("cells",{}):
            if c not in cols: cols.append(c)
    sheet(wb, "KA Matrix", ["Knowledge Area","Scope Note"]+cols, [[r["knowledgeArea"], g(r,"scopeNote")]+[r.get("cells",{}).get(c,"") for c in cols] for r in m["kaMatrix"]], [30,28]+[22]*len(cols))
    sheet(wb, "Rules", ["Rule ID","Name","Rule Scope","Rule Statement","Applies To","Requirement","Status","Source"],
          [[r["id"],r["name"],g(r,"scope"),g(r,"statement"),g(r,"appliesTo"),g(r,"requirement"),S(r),g(r,"origin",src)] for r in m["rulesGov"]], [10,30,14,70,20,12,18,30])
    sheet(wb, "Controls", ["Control ID","Name","Control Type","Applies To","Control Objective / Evaluation","Outcome","Status","Source"],
          [[c["id"],c["name"],g(c,"controlType"),g(c,"appliesTo"),g(c,"objective"),g(c,"outcome"),S(c),g(c,"origin",src)] for c in m["controls"]], [10,34,22,16,60,40,20,30])
    sheet(wb, "Roles", ["Role ID","Role Name","Accountability Type","Responsibility","Applies To","Status","Source"],
          [[r["id"],r["name"],g(r,"accountability"),g(r,"responsibility"),g(r,"appliesTo"),S(r),g(r,"origin",src)] for r in m["roles"]], [10,32,24,60,30,20,30])
    sheet(wb, "Decision Rights", ["Decision Right ID","Name","Holder Role ID","Applies To","Decision Definition","Requirement","Status","Source"],
          [[d["id"],d["name"],g(d,"holder"),g(d,"appliesTo"),g(d,"definition"),g(d,"requirement"),S(d),g(d,"origin",src)] for d in m["decisionRights"]], [12,34,12,16,60,26,20,30])
    sheet(wb, "Exceptions", ["Exception ID","Name","Applies To Transition","Trigger / Basis","Approving Authority","Required Conditions","Lifecycle Status Values","Status","Source"],
          [[x["id"],x["name"],g(x,"transition"),g(x,"basis"),g(x,"authority"),g(x,"conditions"),g(x,"statusValues"),S(x),g(x,"origin",src)] for x in m["exceptions"]], [10,34,16,50,30,50,28,20,30])
    sheet(wb, "Artefacts", ["Artefact ID","Name","Artefact Type","Produced / Modified In","Producing Activity","Evidence Use","Status","Source"],
          [[a["id"],a["name"],g(a,"artefactType"),g(a,"producedIn"),g(a,"producingActivity"),g(a,"evidenceUse"),S(a),g(a,"origin",src)] for a in m["artefacts"]], [10,40,18,18,30,60,20,30])
    sheet(wb, "Evidence", ["Evidence ID","Name","Evidence Type","Relates To","Evidence Description","Requirement","Status","Source"],
          [[e["id"],e["name"],g(e,"evidenceType"),g(e,"relatesTo"),g(e,"description"),g(e,"requirement"),S(e),g(e,"origin",src)] for e in m["evidence"]], [10,36,20,16,60,26,20,30])
    sheet(wb, "Composite States", ["Composite ID","Name","Global State ID","Global Sub-State ID","Included KA State Dimensions","Context / Point in Time","Vector Type","Status","Source"],
          [[c["id"],c["name"],g(c,"globalState"),g(c,"subState"),g(c,"kaDimensions"),g(c,"context"),g(c,"vectorType"),S(c),g(c,"origin",src)] for c in m["compositeStates"]], [10,36,14,16,70,40,18,20,30])
    # Element Register
    reg = []
    for s in m["globalStates"]: reg.append([s["id"],"Global State",s["name"],"",g(s,"definition"),g(s,"semanticClass"),"Global",S(s),g(s,"notes"),g(s,"origin",src)])
    for s in m["subStates"]: reg.append([s["id"],"Sub-State",s["name"],g(s,"parent"),g(s,"definition"),g(s,"semanticClass"),"Global",S(s),g(s,"notes"),g(s,"origin",src)])
    for k, typ, pk, dk, cls in [("activities","Activity","permittedIn","name","Verb / behaviour"),("transitions","Transition","source","guardSummary","Verb / behaviour"),("events","Event","","meaning","Trigger"),("entryConditions","Entry Condition","appliesTo","predicate","Predicate"),("exitConditions","Exit Condition","appliesTo","predicate","Predicate"),("invariants","Invariant","appliesTo","predicate","Predicate"),("guards","Transition Guard","transition","predicate","Predicate"),("kaInteractions","KA Interaction","subState","description","KA state category"),("rulesGov","Rule","appliesTo","statement","Rule"),("controls","Control","appliesTo","objective","Control"),("roles","Role","appliesTo","responsibility","Actor"),("decisionRights","Decision Right","appliesTo","definition","Authority"),("exceptions","Exception","transition","basis","Governed deviation"),("artefacts","Artefact","producedIn","evidenceUse","Artefact"),("evidence","Evidence","relatesTo","description","Proof"),("compositeStates","Composite State","subState","kaDimensions","State vector")]:
        for r in m[k]: reg.append([r["id"],typ,g(r,"name",g(r,"predicate")),g(r,pk) if pk else "",g(r,dk),cls,"Global",S(r),g(r,"notes"),g(r,"origin",src)])
    sheet(wb, "Element Register", ["Element ID","Element Type","Name","Parent / Applies To ID","Definition / Purpose","Semantic Class","Scope","Status","Notes","Source"], reg, [12,18,34,16,60,20,10,20,30,30])
    sheet(wb, "Relationships", ["Relationship ID","Source Element ID","Relationship Type","Target Element ID","Cardinality","Optionality","Rationale / Notes","Status","Source"],
          [[r["id"],r["source"],r["type"],r["target"],g(r,"cardinality"),g(r,"optionality"),g(r,"rationale"),S(r),g(r,"origin",src)] for r in m["relationships"]], [12,16,30,16,12,16,60,20,30])
    dd = [["Metamodel","Global Data State","Fundamental condition of a Data Asset in the Global Data Lifecycle."],["Metamodel","Sub-State","Specialized condition nested in a Global State and inheriting its rules."],["Metamodel","Activity","Permissible verb performed in a State/Sub-State; may be state-preserving."],["Metamodel","Transition","Authorized change between source and target states."],["Metamodel","Event","Occurrence that triggers evaluation of a transition."],["Metamodel","Entry Condition","Predicate that must be satisfied before entry."],["Metamodel","Exit Condition","Predicate that must be satisfied before exit."],["Metamodel","Invariant","Predicate that must remain true throughout State occupancy."],["Metamodel","Transition Guard","Predicate controlling a specific transition."],["Metamodel","KA State","Condition in an independent Knowledge Area state model."],["Metamodel","Rule","Normative statement establishing a constraint."],["Metamodel","Control","Mechanism that evaluates/enforces rules or predicates."],["Metamodel","Decision Right","Authority to approve, reject or condition a transition."],["Metamodel","Exception","Governed, conditional deviation from a rule."],["Metamodel","Artefact","Deliverable produced or modified by an Activity."],["Metamodel","Evidence","Information demonstrating satisfaction of conditions, controls or decisions."],["Metamodel","Composite State","State vector combining Global and applicable KA States at a point in time."],["Column","Status","Proposed / illustrative until formally approved."],["Column","Trace","The context item, placement rule or drafting step a derived row came from."],["QA","State Test","Test whether “The Data Asset is in [name]” meaningfully describes a condition."]]
    sheet(wb, "Data Dictionary", ["Category","Term / Column","Definition / Use"], dd, [12,24,90])
    srcs = m.get("sources") or [{"id":"SRC-001","source":src,"type":"Primary design basis","location":"Private model store","use":"All content","limitations":"Proposed / illustrative status retained."}]
    sheet(wb, "Source Register", ["Source ID","Source","Source Type","Location / Medium","Use in Workbook","Limitations"], [[s.get("id"),s.get("source"),s.get("type"),s.get("location"),s.get("use"),s.get("limitations")] for s in srcs], [10,50,22,30,30,40])
    cc = m.get("contextCapture")
    ccrows = []
    if cc:
        for sec in ["inputs","processes","deliverables","suppliers","participants","consumers","techniques","tools","metrics","goals","businessDrivers"]:
            for i, it in enumerate(cc.get(sec, []), 1):
                if isinstance(it, dict):
                    ccrows.append([sec, it.get("id", f"{sec[:3].upper()}-{i}"), it.get("name",""), it.get("phase",""), "; ".join(it.get("subActivities",[])) if it.get("subActivities") else "", it.get("note","")])
                else: ccrows.append([sec, f"{sec[:3].upper()}-{i}", it, "", "", ""])
    sheet(wb, "Context Capture", ["Section","Item ID","Item","Phase","Sub-activities","Note"], ccrows, [16,10,50,8,70,40], note="Structured capture of the Knowledge Area context diagram this model was derived from. Empty for the Global FTS.")
    sheet(wb, "Derivation Trace", ["Element ID","Element Type","Name / Predicate","Derived From"], [[t["elementId"],t["elementType"],t["name"],t["derivedFrom"]] for t in m.get("derivationTrace",[])], [12,18,50,50], note="Where each derived row came from: a context item, a placement rule, a lifecycle source, or the drafting step.")
    sheet(wb, "QA Findings", ["Severity","Rule","Element","Finding"], [[f["severity"],f["rule"],f["element"],f["finding"]] for f in m.get("qaFindings",[])], [10,10,14,90], note="Soft findings against the retained decisions. Nothing is blocked; each is for your judgement.")
    leg = [["T"+str(i+1), t["id"], t["name"]] for i, t in enumerate(m["transitions"])] + [["E"+str(i+1), e["id"], e["name"]] for i, e in enumerate(m["events"])] + [["G"+str(i+1), x["id"], x["name"]] for i, x in enumerate(m["guards"])] + [["A"+str(i+1), a["id"], a["name"]] for i, a in enumerate(m["activities"])]
    sheet(wb, "Legend", ["Code","Element ID","Name"], leg, [8,14,60], note="Short codes used by the simplified diagram: T transition, E event, G guard, A activity, numbered in model order.")
    wb.save(out)
    return out

if __name__ == "__main__":
    m = json.load(open(sys.argv[1], encoding="utf-8"))
    print("Wrote", export(m, sys.argv[2]))
